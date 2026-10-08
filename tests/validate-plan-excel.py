"""Inspect the XLSX actually downloaded in the browser integration test."""
import json
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

path = '/tmp/action-plan-template.xlsx'
with ZipFile(path) as package:
    for name in package.namelist():
        if name.endswith(('.xml', '.rels')):
            ET.fromstring(package.read(name))
w = load_workbook(path)
assert w.sheetnames == ['Summary', 'Actions']
s = w['Actions']
assert s['B7'].value == 'Ação independente'
assert isinstance(s['D7'].value, datetime)
assert s['D7'].number_format == 'dd/mm/yyyy'
assert s.freeze_panes == 'C7'
assert s.tables['ActionPlan'].ref == 'A6:K8'
assert s['F7'].data_type == 'f' and 'ISNUMBER' in s['F7'].value
assert s['H7'].data_type == 'f' and 'TODAY()' in s['H7'].value
assert len(s.conditional_formatting) == 1
assert len(s.data_validations.dataValidation) == 1
assert s.print_options is not None and s.page_setup.orientation == 'landscape'
assert w['Summary']['A6'].value == "=COUNT('Actions'!A7:A8)"
assert w['Summary']['F10'].number_format == '0%'
assert len(w['Summary'].merged_cells.ranges) > 10
cached = load_workbook(path, data_only=True)
assert cached['Summary']['A6'].value == 2
assert cached['Actions']['F7'].value == 5
assert load_workbook('/tmp/action-plan-safe-text.xlsx')['Actions']['B7'].data_type == 's'
assert load_workbook('/tmp/action-plan-safe-text.xlsx')['Actions']['B7'].value.startswith('=HYPERLINK')
for filename in ['/tmp/action-plan-export.json', '/tmp/action-plan-backup.json']:
    payload = json.loads(Path(filename).read_text())
    assert payload['type'] == 'action-plans' and 'projects' not in payload
print('PASS: Excel package/XML, native dates, formulas and caches, dashboard, table/filter/freeze, styles, validation, print layout, safe literal text and scoped JSON artifacts')
